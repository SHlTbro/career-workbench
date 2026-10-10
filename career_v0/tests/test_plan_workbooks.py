from __future__ import annotations

import copy
import io
import unittest

from openpyxl import load_workbook

from career_v0.plan_workbooks import export_plan_workbook, import_plan_workbook
from career_v0.tests.test_plans import INVENTORY, plan
from career_v0.workbooks import WorkbookError


def fixture():
    first = plan()
    second = copy.deepcopy(first)
    second["plan_id"] = "PLAN-JOB-2"
    second["job_ref"]["job_id"] = "JOB-2"
    inventory = [{**item, "purpose": "合成用途", "core_stack": ["Python"]} for item in INVENTORY]
    entries = [
        {"plan": first, "job": {"job_id": "JOB-1", "company": "公司一", "title": "岗位一"},
         "research": {"jd": {"source_url": "https://example.org/1"}, "business_summary": "业务一"}},
        {"plan": second, "job": {"job_id": "JOB-2", "company": "公司二", "title": "岗位二"},
         "research": {"jd": {"source_url": "https://example.org/2"}, "business_summary": "业务二"}},
    ]
    data, snapshot = export_plan_workbook("CMP-1", "BAT-1", entries, inventory)
    current = {plan_id: {key: row[key] for key in ("plan_revision", "input_snapshot_sha256", "plan_content_sha256")}
               for plan_id, row in snapshot["rows"].items()}
    return data, snapshot, current


def save(wb) -> bytes:
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


class PlanWorkbookTests(unittest.TestCase):
    def test_two_plan_decisions_round_trip_only_to_preview(self):
        data, snapshot, current = fixture()
        wb = load_workbook(io.BytesIO(data))
        self.assertEqual(wb.sheetnames, ["使用说明", "方案总览", "逐项目取舍", "来源与差距", "_meta"])
        main = wb["方案总览"]
        main.cell(2, 1).value = "APPROVE"
        main.cell(3, 1).value = "REQUEST_CHANGES"
        main.cell(3, 2).value = "请收紧业务表述"
        preview = import_plan_workbook(save(wb), snapshot, current)
        self.assertEqual(preview["status"], "PENDING_CONFIRMATION")
        self.assertEqual([row["decision"] for row in preview["valid_rows"]], ["APPROVE", "REQUEST_CHANGES"])
        self.assertEqual(preview["conflicts"], [])

    def test_one_stale_plan_is_isolated(self):
        data, snapshot, current = fixture()
        wb = load_workbook(io.BytesIO(data))
        wb["方案总览"].cell(2, 1).value = "APPROVE"
        wb["方案总览"].cell(3, 1).value = "APPROVE"
        current["PLAN-JOB-2"]["plan_revision"] = 2
        preview = import_plan_workbook(save(wb), snapshot, current)
        self.assertEqual([row["plan_id"] for row in preview["valid_rows"]], ["PLAN-JOB-1"])
        self.assertEqual(preview["conflicts"], [{"plan_id": "PLAN-JOB-2", "reason": "STALE_PLAN_OR_INPUT"}])

    def test_approval_with_change_request_is_conflict(self):
        data, snapshot, current = fixture()
        wb = load_workbook(io.BytesIO(data))
        wb["方案总览"].cell(2, 1).value = "APPROVE"
        wb["方案总览"].cell(2, 2).value = "把项目改成另一种策略"
        preview = import_plan_workbook(save(wb), snapshot, current)
        self.assertEqual(preview["conflicts"], [{"plan_id": "PLAN-JOB-1", "reason": "APPROVAL_WITH_CHANGE_REQUEST"}])
        wb = load_workbook(io.BytesIO(data))
        wb["方案总览"].cell(2, 1).value = "REQUEST_CHANGES"
        preview = import_plan_workbook(save(wb), snapshot, current)
        self.assertEqual(preview["conflicts"], [{"plan_id": "PLAN-JOB-1", "reason": "CHANGE_REQUEST_NEEDS_COMMENT"}])

    def test_readonly_formula_and_invalid_decision_rejected(self):
        data, snapshot, current = fixture()
        wb = load_workbook(io.BytesIO(data))
        wb["逐项目取舍"].cell(2, 4).value = "DROP"
        with self.assertRaisesRegex(WorkbookError, "read-only conflict"):
            import_plan_workbook(save(wb), snapshot, current)
        wb = load_workbook(io.BytesIO(data))
        wb["来源与差距"].cell(2, 2).value = "=1+1"
        with self.assertRaisesRegex(WorkbookError, "formula"):
            import_plan_workbook(save(wb), snapshot, current)
        wb = load_workbook(io.BytesIO(data))
        wb["方案总览"].cell(2, 1).value = "YES"
        with self.assertRaisesRegex(WorkbookError, "invalid decision"):
            import_plan_workbook(save(wb), snapshot, current)


if __name__ == "__main__":
    unittest.main()
