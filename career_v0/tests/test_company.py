from __future__ import annotations

import io
import json
import tempfile
import unittest
from pathlib import Path
from subprocess import CompletedProcess
from unittest.mock import patch

from openpyxl import load_workbook

from career_v0.company_runner import prepare_company_dossier, search_company_sources
from career_v0.bridge import _source_summary
from career_v0.discovery import DiscoveryError
from career_v0.research import ResearchError, attach_company_dossier
from career_v0.service import CareerService
from career_v0.tests.test_m1 import SYNTHETIC_MASTER, job, preferences
from career_v0.tests.test_research import JD, dossier


def trace(events):
    return ("\n".join(json.dumps(row, ensure_ascii=False) for row in events) + "\n").encode()


class CompanyResearchTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        master = self.root / "master.md"
        master.write_bytes(SYNTHETIC_MASTER)
        self.dsh = self.root / "dsh.cmd"
        self.dsh.write_text("@echo off", encoding="utf-8")
        self.service = CareerService(self.root / "state", master_path=master, allow_simulated_owner=True)
        self.service.create_campaign(preferences(), campaign_id="CMP-COMPANY")
        self.service.confirm_campaign("CMP-COMPANY", actor_channel="SIMULATED_OWNER_TEST")
        second = job(2)
        second["company"] = "公司1"
        self.batch = self.service.create_batch("CMP-COMPANY", [job(1), second], batch_id="BAT-COMPANY")
        path, export = self.service.export_candidates("BAT-COMPANY")
        workbook = load_workbook(path)
        for row in (2, 3):
            workbook["岗位候选"].cell(row, 1).value = "TRUE"
        stream = io.BytesIO()
        workbook.save(stream)
        preview = self.service.preview_selection(export["export_id"], stream.getvalue())
        self.service.confirm_selection(preview["preview_id"], actor_channel="SIMULATED_OWNER_TEST",
                                       accept_shortfall=True)
        self.job_ids = [item["job_id"] for item in self.service.selected_jobs("BAT-COMPANY")]

    def tearDown(self):
        self.tmp.cleanup()

    def source(self):
        url = "https://company.example/about"
        search = self.service.ingest_discovery_trace("CMP-COMPANY", trace([
            {"type": "tool_call", "callId": "s1", "tool": "web_search", "input": {"queries": ["公司1 官网"]}},
            {"type": "tool_result", "callId": "s1", "status": "completed",
             "result": f"Sources:\n- [公司1 公开业务](" + url + ")"},
        ]), purpose="COMPANY_RESEARCH", company="公司1")
        body = ("公司1 是一家公开介绍其业务的企业。公司1 公开页面说明其产品包括智能协作平台和开发者工具，"
                "服务于企业知识工作场景，并提供检索与流程编排能力。") * 3
        discovery = self.service.extend_discovery_trace(search["discovery_id"], trace([
            {"type": "tool_call", "callId": "f1", "tool": "web_fetch", "input": {"url": url}},
            {"type": "tool_result", "callId": "f1", "status": "completed",
             "result": f"Fetched {url} (HTTP 200)\n\nExternal web content follows.\n\n{body}"},
        ]))
        return discovery, discovery["captures"][0], body

    def test_company_source_is_selected_only_and_cannot_create_job_batch(self):
        discovery, capture, _ = self.source()
        summary = _source_summary(discovery)
        self.assertNotIn("raw_result", summary["captures"][0])
        self.assertNotIn("content_text", summary["captures"][0])
        with self.assertRaisesRegex(ValueError, "company research discovery"):
            self.service.create_batch("CMP-COMPANY", [job(3)], discovery_id=discovery["discovery_id"])
        with self.assertRaisesRegex(ResearchError, "independent readable"):
            self.service.record_company_dossier("BAT-COMPANY", self.job_ids[0], discovery["discovery_id"],
                                                "missing", "公司1", "公司1 公开页面说明其产品包括智能协作平台和开发者工具")
        with self.assertRaisesRegex(ValueError, "G1 selection"):
            self.service.record_company_dossier("BAT-COMPANY", "J-NOT-SELECTED", discovery["discovery_id"],
                                                capture["capture_id"], "公司1", "公司1 公开页面说明其产品包括智能协作平台和开发者工具")

    def test_exact_quote_and_same_company_reuse(self):
        discovery, capture, _ = self.source()
        identity = "公司1 是一家公开介绍其业务的企业。"
        business = "公司1 公开页面说明其产品包括智能协作平台和开发者工具，服务于企业知识工作场景"
        with self.assertRaisesRegex(ResearchError, "quotes must occur exactly"):
            self.service.record_company_dossier("BAT-COMPANY", self.job_ids[0], discovery["discovery_id"],
                                                capture["capture_id"], identity, "不存在的公司业务摘录")
        company = self.service.record_company_dossier("BAT-COMPANY", self.job_ids[0], discovery["discovery_id"],
                                                      capture["capture_id"], identity, business)
        self.assertEqual(company["revision"], 1)
        self.assertEqual(company["source_authority"], "NOT_ASSESSED")
        role_dossier = {"company": "公司1", "business_evidence": [],
                        "unknowns": ["公司业务线未独立核实"], "jd": {"job_id": self.job_ids[1]}}
        reused = attach_company_dossier(role_dossier, company)
        self.assertEqual(reused["jd"]["job_id"], self.job_ids[1])
        self.assertEqual(reused["company_dossier_id"], company["company_dossier_id"])
        self.assertIn("公开来源权威性与时效未核实", reused["unknowns"])
        with self.assertRaisesRegex(ResearchError, "identity differs"):
            attach_company_dossier({**role_dossier, "company": "另一家公司"}, company)
        for selected in self.service.selected_jobs("BAT-COMPANY"):
            self.service.confirm_owner_jd("BAT-COMPANY", selected["job_id"], JD,
                                          actor_channel="SIMULATED_OWNER_TEST")
            self.service.save_research_dossier("BAT-COMPANY", dossier(selected, full=True))
        refreshed = [self.service.refresh_research_with_company("BAT-COMPANY", job_id)
                     for job_id in self.job_ids]
        self.assertEqual([row["revision"] for row in refreshed], [2, 2])
        self.assertEqual({row["company_dossier_id"] for row in refreshed}, {company["company_dossier_id"]})
        self.assertEqual([row["jd"]["job_id"] for row in refreshed], self.job_ids)

    def test_model_extracts_only_exact_quotes(self):
        discovery, capture, _ = self.source()
        proposed = {"identity_quote": "公司1 是一家公开介绍其业务的企业。",
                    "business_quote": "公司1 公开页面说明其产品包括智能协作平台和开发者工具，服务于企业知识工作场景"}
        output = trace([{"type": "final", "text": json.dumps(proposed, ensure_ascii=False)}])
        with patch("career_v0.company_runner.subprocess.run",
                   return_value=CompletedProcess([], 0, stdout=output, stderr=b"")):
            result = prepare_company_dossier(self.service, "BAT-COMPANY", self.job_ids[0],
                                             discovery["discovery_id"], capture["capture_id"],
                                             dsh_bin=self.dsh)
        self.assertEqual(result["company_dossier"]["business_quote"], proposed["business_quote"])
        self.assertIsNone(result["research"])

    def test_company_search_rejects_non_search_tools(self):
        with patch("career_v0.company_runner.subprocess.run",
                   return_value=CompletedProcess([], 0, stdout=trace([
                       {"type": "tool_call", "callId": "x", "tool": "file_write", "input": {}},
                   ]), stderr=b"")):
            with self.assertRaisesRegex(DiscoveryError, "unexpected tools"):
                search_company_sources(self.service, "BAT-COMPANY", self.job_ids[0], dsh_bin=self.dsh)


if __name__ == "__main__":
    unittest.main()
