from __future__ import annotations

import io
import os
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from career_v0.domain import DomainError
from career_v0.inventory import candidate_inventory
from career_v0.service import CareerService
from career_v0.workbooks import WorkbookError


SYNTHETIC_MASTER = ("# 合成候选人\nsample@example.invalid ｜ 合成联系方式\n"
                    "## 求职亮点\n- 长期使用 10+ AI 平台，完成 7+ AI 项目，保留开发过程留痕。\n"
                    "## 教育背景\n合成大学｜计算机专业｜2027 届\n"
                    "**实习经历｜合成团队｜工程实践**\n- 验证合成工作流的来源与输出。\n"
                    "## 核心能力\nPython、Agent 与检索工程。\n"
                    "# 项目经历\n### Edu Agent —— 教育平台\n- 内容\n### YY Workflow —— 工作流\n- 内容\n"
                    "###Local Codebase MCP —— 本地网关\n- 内容\n### JevTest Triage —— 分诊\n- 内容\n"
                    "### AI Engineering Knowledge Base —— 知识库\n**技术雷达**\n- 内容\n"
                    "# 其他项目\n### Career Agent：求职平台\n### ToSign Studio：字体平台\n").encode("utf-8")


def confirmed_inventory() -> dict:
    candidate = candidate_inventory(SYNTHETIC_MASTER)
    for item in candidate["projects"]:
        item.update({"purpose": "合成测试用途", "core_stack": ["Python"],
                     "status": "IN_PROGRESS", "include_in_resume": True})
        for child in item["child_candidates"]:
            child["owner_decision"] = "SUBPROJECT"
    return candidate


def preferences() -> dict:
    return {
        "recruitment_type": "AUTUMN_CAMPUS", "allow_internships": False,
        "graduation_year": 2027, "target_roles": ["AI Agent 开发"],
        "city_scope": {"mode": "ANY", "values": []}, "company_scope": "ANY",
        "adjacent_role_policy": {"frontend": False, "algorithm": False, "sales": False},
        "salary": {"currency": "CNY", "basis": "PRE_TAX_BASE_MONTHLY",
                   "minimum": 10000, "comparison": "GTE", "unknown_policy": "REVIEW_BUCKET"},
        "batch_target_count": 50,
        "project_inventory": confirmed_inventory(),
        "identity_signals": {
            "ID-PLATFORMS": {"source_quote": "10+ AI 平台", "approved_expression": "长期使用 10+ AI 平台", "owner_confirmed": True},
            "ID-PROJECTS": {"source_quote": "7+ AI 项目", "approved_expression": "完成 7+ AI 项目", "owner_confirmed": True},
            "ID-PROVENANCE": {"source_quote": "开发过程留痕", "approved_expression": "保留开发过程留痕", "owner_confirmed": True},
        },
    }


def job(number: int, *, salary: int | None = 12000, recruitment: str = "AUTUMN_CAMPUS", source_url: str | None = None) -> dict:
    return {
        "platform_job_id": f"J{number}", "source_platform": "official",
        "company": f"公司{number}", "title": "AI Agent 工程师", "city": "郑州",
        "salary_raw": f"{salary}元/月" if salary is not None else None,
        "salary_min_monthly": salary, "salary_basis": "PRE_TAX_BASE_MONTHLY" if salary is not None else None,
        "recruitment_type": recruitment, "graduation_years": [2027],
        "required_degree": "BACHELOR", "required_experience_years": 0,
        "role_match": "MATCH", "role_match_evidence": "AI Agent 工程师",
        "role_family": "AI_AGENT_DEV", "opening_status": "OPEN",
        "source_url": source_url or f"https://example.org/jobs/{number}",
        "jd_completeness": "FULL", "application_status": "NOT_FOUND",
    }


def save(wb) -> bytes:
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


class M1RoundTripTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.master_path = os.path.join(self.tmp.name, "master.md")
        Path(self.master_path).write_bytes(SYNTHETIC_MASTER)
        self.service = CareerService(self.tmp.name, master_path=self.master_path, allow_simulated_owner=True)
        self.campaign = self.service.create_campaign(preferences(), campaign_id="CMP-TEST")

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def _batch(self):
        self.campaign = self.service.confirm_campaign("CMP-TEST", actor_channel="SIMULATED_OWNER_TEST")
        return self.service.create_batch("CMP-TEST", [job(1), job(2), job(3, salary=None), job(4, recruitment="INTERNSHIP")], batch_id="BAT-TEST")

    def test_g0_blocks_discovery_and_research(self):
        with self.assertRaisesRegex(DomainError, "G0"):
            self.service.create_batch("CMP-TEST", [job(1)], batch_id="BAT-TEST")
        with self.assertRaisesRegex(DomainError, "G1"):
            self.service.selected_jobs("BAT-TEST")
        with self.assertRaisesRegex(DomainError, "Owner action"):
            self.service.confirm_campaign("CMP-TEST", actor_channel="MODEL_OUTPUT")
        with self.assertRaisesRegex(DomainError, "Owner action"):
            self.service.confirm_campaign("CMP-TEST", actor_channel="OWNER_UI:unverified")

    def test_g0_draft_can_be_saved_incrementally_but_not_confirmed(self):
        self.service.create_campaign({}, campaign_id="CMP-DRAFT")
        revised = self.service.revise_campaign("CMP-DRAFT", {"graduation_year": 2027},
                                               actor_channel="SIMULATED_OWNER_TEST")
        self.assertEqual(revised["status"], "PREFERENCE_DRAFT")
        self.assertEqual(revised["revision"], 2)
        with self.assertRaises(DomainError):
            self.service.confirm_campaign("CMP-DRAFT", actor_channel="SIMULATED_OWNER_TEST")
        completed = self.service.revise_campaign("CMP-DRAFT", preferences(),
                                                 actor_channel="SIMULATED_OWNER_TEST")
        self.assertEqual(completed["status"], "PREFERENCE_DRAFT")
        self.assertEqual(self.service.confirm_campaign("CMP-DRAFT", actor_channel="SIMULATED_OWNER_TEST")["status"],
                         "PREFERENCE_CONFIRMED")

    def test_g0_inventory_requires_explicit_project_and_subproject_choices(self):
        incomplete = preferences()
        incomplete["project_inventory"] = candidate_inventory(SYNTHETIC_MASTER)
        self.service.create_campaign(incomplete, campaign_id="CMP-INCOMPLETE")
        with self.assertRaisesRegex(DomainError, "explicit include"):
            self.service.confirm_campaign("CMP-INCOMPLETE", actor_channel="SIMULATED_OWNER_TEST")
        approved = preferences()
        approved["project_inventory"]["projects"][4]["child_candidates"][0]["owner_decision"] = None
        self.service.create_campaign(approved, campaign_id="CMP-CHILD")
        with self.assertRaisesRegex(DomainError, "nested candidates need Owner decisions"):
            self.service.confirm_campaign("CMP-CHILD", actor_channel="SIMULATED_OWNER_TEST")
        missing_signal = preferences()
        missing_signal["identity_signals"]["ID-PROVENANCE"]["owner_confirmed"] = False
        self.service.create_campaign(missing_signal, campaign_id="CMP-SIGNAL")
        with self.assertRaisesRegex(DomainError, "explicit Owner confirmation"):
            self.service.confirm_campaign("CMP-SIGNAL", actor_channel="SIMULATED_OWNER_TEST")
        excluded_without_metadata = preferences()
        excluded_without_metadata["project_inventory"]["projects"][0]["include_in_resume"] = False
        excluded_without_metadata["project_inventory"]["projects"][0]["purpose"] = None
        self.service.create_campaign(excluded_without_metadata, campaign_id="CMP-METADATA")
        with self.assertRaisesRegex(DomainError, "purpose missing"):
            self.service.confirm_campaign("CMP-METADATA", actor_channel="SIMULATED_OWNER_TEST")

    def test_three_pools_and_shortfall(self):
        batch = self._batch()
        self.assertEqual(batch["counts"], {"qualified": 2, "unknown": 1, "excluded": 1})
        self.assertEqual(batch["target_shortfall"], 48)
        self.assertEqual({j["pool"] for j in batch["jobs"].values()}, {"qualified", "unknown", "excluded"})

    def test_unknown_role_experience_jd_and_opening_never_count_as_qualified(self):
        self.campaign = self.service.confirm_campaign("CMP-TEST", actor_channel="SIMULATED_OWNER_TEST")
        cases = []
        for number, missing in enumerate(("role_match", "required_experience_years", "jd_completeness", "opening_status"), 1):
            row = job(number)
            row.pop(missing)
            cases.append(row)
        batch = self.service.create_batch("CMP-TEST", cases, batch_id="BAT-UNKNOWN")
        self.assertEqual(batch["counts"], {"qualified": 0, "unknown": 4, "excluded": 0})

    def test_adjacent_roles_and_company_filter(self):
        self.campaign = self.service.confirm_campaign("CMP-TEST", actor_channel="SIMULATED_OWNER_TEST")
        frontend = job(1)
        frontend["role_family"] = "FRONTEND"
        batch = self.service.create_batch("CMP-TEST", [frontend], batch_id="BAT-FRONTEND")
        self.assertEqual(batch["counts"]["excluded"], 1)
        filtered = preferences()
        filtered["company_scope"] = "FILTERED"
        filtered["company_filters"] = {"names": ["指定公司"]}
        self.service.create_campaign(filtered, campaign_id="CMP-FILTERED")
        self.service.confirm_campaign("CMP-FILTERED", actor_channel="SIMULATED_OWNER_TEST")
        another = self.service.create_batch("CMP-FILTERED", [job(2)], batch_id="BAT-FILTERED")
        self.assertEqual(another["counts"]["excluded"], 1)

    def test_confirmed_salary_unknown_policy_excludes_unknown_salary(self):
        selected_policy = preferences()
        selected_policy["salary"]["unknown_policy"] = "EXCLUDE"
        self.service.create_campaign(selected_policy, campaign_id="CMP-EXCLUDE")
        self.service.confirm_campaign("CMP-EXCLUDE", actor_channel="SIMULATED_OWNER_TEST")
        batch = self.service.create_batch("CMP-EXCLUDE", [job(5, salary=None)], batch_id="BAT-SALARY")
        self.assertEqual(batch["counts"], {"qualified": 0, "unknown": 0, "excluded": 1})

    def test_campaign_revision_persists_and_stales_old_batch(self):
        batch = self._batch()
        restarted = CareerService(self.tmp.name, master_path=self.master_path, allow_simulated_owner=True)
        self.assertEqual(restarted.store.get("batches", batch["batch_id"])["counts"]["qualified"], 2)
        revised_prefs = preferences()
        revised_prefs["city_scope"] = {"mode": "CITIES", "values": ["北京"]}
        revised = restarted.revise_campaign("CMP-TEST", revised_prefs, actor_channel="SIMULATED_OWNER_TEST")
        self.assertEqual(revised["status"], "PREFERENCE_DRAFT")
        self.assertEqual(revised["affected_batch_ids"], ["BAT-TEST"])
        self.assertEqual(restarted.store.get_revision("campaigns", "CMP-TEST", 2)["city_scope"]["mode"], "ANY")
        with self.assertRaisesRegex(DomainError, "G0"):
            restarted.create_batch("CMP-TEST", [job(8)], batch_id="BAT-NEW")
        restarted.confirm_campaign("CMP-TEST", actor_channel="SIMULATED_OWNER_TEST")
        with self.assertRaisesRegex(DomainError, "stale"):
            restarted.export_candidates("BAT-TEST")

    def test_sorted_rows_stable_id_and_confirmed_g1(self):
        batch = self._batch()
        path, snapshot = self.service.export_candidates(batch["batch_id"])
        self.assertTrue(path.exists())
        wb = load_workbook(path)
        self.assertEqual(wb.sheetnames, ["使用说明", "岗位候选", "待核实", "排除摘要", "_meta"])
        ws = wb["岗位候选"]
        rows = [[cell.value for cell in row] for row in ws.iter_rows(min_row=2)]
        self.assertEqual(len(rows), 2)
        for col, value in enumerate(rows[1], 1):
            ws.cell(2, col).value = value
        for col, value in enumerate(rows[0], 1):
            ws.cell(3, col).value = value
        ws.cell(2, 1).value = "TRUE"
        chosen_id = ws.cell(2, 4).value
        data = save(wb)
        preview = self.service.preview_selection(snapshot["export_id"], data)
        self.assertEqual(preview["selected_job_ids"], [chosen_id])
        with self.assertRaisesRegex(DomainError, "Owner action"):
            self.service.confirm_selection(preview["preview_id"], actor_channel="MODEL_OUTPUT")
        with self.assertRaisesRegex(DomainError, "smaller qualified batch"):
            self.service.confirm_selection(preview["preview_id"], actor_channel="SIMULATED_OWNER_TEST")
        selection = self.service.confirm_selection(preview["preview_id"], actor_channel="SIMULATED_OWNER_TEST", accept_shortfall=True)
        self.assertEqual([j["job_id"] for j in self.service.selected_jobs("BAT-TEST")], [chosen_id])
        self.assertEqual(selection["confirmed_by"], "SIMULATED_OWNER")
        second_preview = self.service.preview_selection(snapshot["export_id"], data)
        repeated = self.service.confirm_selection(second_preview["preview_id"], actor_channel="SIMULATED_OWNER_TEST", accept_shortfall=True)
        self.assertEqual(repeated["selection_id"], selection["selection_id"])

    def test_empty_and_illegal_selections_do_not_confirm(self):
        batch = self._batch()
        path, snapshot = self.service.export_candidates(batch["batch_id"])
        data = path.read_bytes()
        preview = self.service.preview_selection(snapshot["export_id"], data)
        with self.assertRaisesRegex(DomainError, "empty selection"):
            self.service.confirm_selection(preview["preview_id"], actor_channel="SIMULATED_OWNER_TEST")
        wb = load_workbook(io.BytesIO(data))
        wb["岗位候选"].cell(2, 1).value = "YES"
        with self.assertRaisesRegex(WorkbookError, "selected"):
            self.service.preview_selection(snapshot["export_id"], save(wb))
        self.assertIsNone(self.service.store.get("selections", "BAT-TEST"))

    def test_readonly_formula_and_identity_conflicts(self):
        batch = self._batch()
        path, snapshot = self.service.export_candidates(batch["batch_id"])
        wb = load_workbook(path)
        wb["岗位候选"].cell(2, 5).value = "伪造公司"
        with self.assertRaisesRegex(WorkbookError, "read-only conflict"):
            self.service.preview_selection(snapshot["export_id"], save(wb))
        wb = load_workbook(path)
        wb["岗位候选"].cell(2, 3).value = "=1+1"
        with self.assertRaisesRegex(WorkbookError, "formula"):
            self.service.preview_selection(snapshot["export_id"], save(wb))
        wb = load_workbook(path)
        wb["岗位候选"].cell(3, 4).value = wb["岗位候选"].cell(2, 4).value
        with self.assertRaisesRegex(WorkbookError, "duplicate job_id"):
            self.service.preview_selection(snapshot["export_id"], save(wb))
        wb = load_workbook(path)
        wb["_meta"].cell(3, 2).value = "CMP-OTHER"
        with self.assertRaisesRegex(WorkbookError, "campaign_id"):
            self.service.preview_selection(snapshot["export_id"], save(wb))
        wb = load_workbook(path)
        wb["待核实"].cell(2, 2).value = "伪造公司"
        with self.assertRaisesRegex(WorkbookError, "read-only conflict"):
            self.service.preview_selection(snapshot["export_id"], save(wb))
        wb = load_workbook(path)
        wb["使用说明"].cell(10, 2).value = "=1+1"
        with self.assertRaisesRegex(WorkbookError, "formula"):
            self.service.preview_selection(snapshot["export_id"], save(wb))

    def test_export_source_formula_is_text(self):
        self.campaign = self.service.confirm_campaign("CMP-TEST", actor_channel="SIMULATED_OWNER_TEST")
        batch = self.service.create_batch("CMP-TEST", [job(1, source_url="=HYPERLINK(\"bad\",\"click\")")], batch_id="BAT-FORMULA")
        path, _ = self.service.export_candidates(batch["batch_id"])
        ws = load_workbook(path)["岗位候选"]
        headers = [c.value for c in ws[1]]
        jd_cell = ws.cell(2, headers.index("jd_url") + 1)
        self.assertEqual(jd_cell.data_type, "s")


if __name__ == "__main__":
    unittest.main()
