from __future__ import annotations

import shutil
import hashlib
import io
import os
import tempfile
import unittest
import copy
import json
import zipfile
from pathlib import Path
from subprocess import CompletedProcess
from unittest.mock import patch

from openpyxl import load_workbook
from fpdf import FPDF
import pypdfium2 as pdfium

from career_v0.domain import DomainError
from career_v0.research import ResearchError, validate_application_card
from career_v0.service import CareerService
from career_v0.planning_runner import PlanningError, prepare_tailoring_plan
from career_v0.tests.test_m1 import SYNTHETIC_MASTER, job, preferences
from career_v0.tests.test_discovery_runner import event_trace, full_fetch_trace, normalized_json
from career_v0.tests.test_plans import plan as plan_fixture


JD = "岗位职责：设计与实现 AI Agent 工作流，建设检索、工具调用与评估链路。任职要求：熟悉 Python、RAG、MCP，2027 届本科毕业生，愿意开展工程验证与团队协作。"


def card(selected_job):
    reasons = {field: "公开来源未找到或尚未到 G2" for field in (
        "original_publisher", "apply_method", "apply_url", "login_required", "deadline",
        "contact_person", "contact_email", "contact_scope", "recommended_resume_ref",
        "last_verified_at")}
    return {"job_id": selected_job["job_id"], "company": selected_job["company"],
            "title": selected_job["title"], "source_url": selected_job["source_url"],
            "source_platform": selected_job["source_platform"], "original_publisher": None,
            "apply_method": None, "apply_url": None, "application_status": "NOT_FOUND",
            "login_required": None, "deadline": None, "opening_status": "OPEN",
            "contact_person": None, "contact_email": None, "contact_scope": None,
            "required_materials": [], "recommended_resume_ref": None,
            "last_verified_at": None, "evidence_refs": [], "next_action": "Owner 查找官方渠道",
            "unknown_reasons": reasons}


def dossier(selected_job, *, full: bool):
    return {"schema": "research-dossier@1", "job_id": selected_job["job_id"],
            "company": selected_job["company"], "business_summary": "合成测试业务摘要",
            "jd": {"job_id": selected_job["job_id"], "revision": 1,
                   "body": JD if full else None,
                   "sha256": hashlib.sha256(JD.encode()).hexdigest() if full else None,
                   "completeness": "FULL" if full else "PARTIAL",
                   "source_url": selected_job["source_url"],
                   "requirement_ids": ["JD-R1"] if full else [],
                   "requirement_quotes": {"JD-R1": "任职要求：熟悉 Python、RAG、MCP，2027 届本科毕业生，愿意开展工程验证与团队协作。"} if full else {}},
            "business_evidence": [{"evidence_id": "COMPANY-1", "status": "INFERRED",
                                   "reasoning": "仅从合成 JD 推断工作范围，未独立核实公司业务"}],
            "fde_map": [{"fde_id": "FDE-1", "status": "INFERRED",
                         **{field: "待核实的合成映射" for field in (
                             "business_scenario", "users_and_team", "workflow", "pain_hypothesis",
                             "ai_entry_point", "engineering_constraints", "jd_capability",
                             "owner_project_evidence", "tailoring_action")}}],
            "application_card": card(selected_job), "unknowns": ["投递渠道未找到"]}


class SelectedResearchTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.master = os.path.join(self.tmp.name, "master.md")
        Path(self.master).write_bytes(SYNTHETIC_MASTER)
        self.service = CareerService(self.tmp.name, master_path=self.master, allow_simulated_owner=True)
        self.service.create_campaign(preferences(), campaign_id="CMP-RESEARCH")
        self.service.confirm_campaign("CMP-RESEARCH", actor_channel="SIMULATED_OWNER_TEST")
        self.batch = self.service.create_batch("CMP-RESEARCH", [job(1), job(2)], batch_id="BAT-RESEARCH")
        self.job_ids = list(self.batch["jobs"])

    def tearDown(self):
        self.tmp.cleanup()

    def select_first(self):
        path, export = self.service.export_candidates("BAT-RESEARCH")
        wb = load_workbook(path)
        ws = wb["岗位候选"]
        ws.cell(2, 1).value = "TRUE"
        selected_id = ws.cell(2, 4).value
        stream = io.BytesIO()
        wb.save(stream)
        preview = self.service.preview_selection(export["export_id"], stream.getvalue())
        self.service.confirm_selection(preview["preview_id"], actor_channel="SIMULATED_OWNER_TEST", accept_shortfall=True)
        return selected_id

    def test_unselected_research_and_owner_jd_are_blocked(self):
        first = self.batch["jobs"][self.job_ids[0]]
        with self.assertRaisesRegex(DomainError, "G1"):
            self.service.save_research_dossier("BAT-RESEARCH", dossier(first, full=False))
        selected_id = self.select_first()
        other = next(item for item in self.batch["jobs"].values() if item["job_id"] != selected_id)
        with self.assertRaisesRegex(DomainError, "G1 selection"):
            self.service.save_research_dossier("BAT-RESEARCH", dossier(other, full=False))

    def test_full_jd_requires_source_capture_or_owner_confirmation(self):
        selected_id = self.select_first()
        selected = self.batch["jobs"][selected_id]
        blocked = self.service.save_research_dossier("BAT-RESEARCH", dossier(selected, full=False))
        self.assertEqual(blocked["status"], "RESEARCH_BLOCKED")
        with self.assertRaisesRegex(ResearchError, "FULL JD requires"):
            self.service.save_research_dossier("BAT-RESEARCH", dossier(selected, full=True))
        with self.assertRaisesRegex(DomainError, "Owner action"):
            self.service.confirm_owner_jd("BAT-RESEARCH", selected_id, JD, actor_channel="MODEL_OUTPUT")
        self.service.confirm_owner_jd("BAT-RESEARCH", selected_id, JD, actor_channel="SIMULATED_OWNER_TEST")
        ready = self.service.save_research_dossier("BAT-RESEARCH", dossier(selected, full=True))
        self.assertEqual(ready["status"], "READY_FOR_PLAN")
        self.assertTrue(ready["jd"]["owner_confirmed"])
        self.assertEqual(ready["revision"], 2)

    def test_application_card_unknown_channel_cannot_claim_ready(self):
        selected = self.batch["jobs"][self.job_ids[0]]
        actual = card(selected)
        actual["application_status"] = "READY_TO_APPLY"
        with self.assertRaisesRegex(ResearchError, "READY_TO_APPLY"):
            validate_application_card(actual, selected)

    def test_verified_company_claim_needs_matching_fetch_quote(self):
        selected_id = self.select_first()
        selected = self.batch["jobs"][selected_id]
        self.service.confirm_owner_jd("BAT-RESEARCH", selected_id, JD, actor_channel="SIMULATED_OWNER_TEST")
        proposed = dossier(selected, full=True)
        proposed["business_evidence"] = [{"evidence_id": "COMPANY-1", "status": "VERIFIED",
                                          "source_url": selected["source_url"],
                                          "source_quote": "this company does pharmaceutical AI"}]
        with self.assertRaisesRegex(ResearchError, "not grounded"):
            self.service.save_research_dossier("BAT-RESEARCH", proposed)

    def ready_plan(self):
        selected_id = self.select_first()
        selected = self.batch["jobs"][selected_id]
        self.service.confirm_owner_jd("BAT-RESEARCH", selected_id, JD, actor_channel="SIMULATED_OWNER_TEST")
        research = self.service.save_research_dossier("BAT-RESEARCH", dossier(selected, full=True))
        campaign = self.service.store.get("campaigns", "CMP-RESEARCH")
        tailored = copy.deepcopy(plan_fixture())
        tailored["plan_id"] = "PLAN-RESEARCH-1"
        tailored["campaign_id"] = "CMP-RESEARCH"
        tailored["job_ref"]["job_id"] = selected_id
        tailored["input_refs"].update({
            "campaign_revision": campaign["revision"], "jd_sha256": research["jd"]["sha256"],
            "master_sha256": campaign["master_sha256"], "research_revision": research["revision"],
            "resume_profile_revision": campaign["resume_profile"]["revision"],
            "project_inventory_revision": campaign["resume_profile"]["project_inventory_revision"],
        })
        tailored["strategy"]["jd_requirement_refs"] = ["JD-R1"]
        required = [row for row in campaign["project_inventory"]["projects"] if row["include_in_resume"]]
        for index, row in enumerate(tailored["project_plan"]):
            source = required[index]
            row["project_id"] = source["project_id"]
            row["source_order"] = source["original_index"]
            row["owner_evidence_refs"] = ["MASTER-" + source["project_id"]]
        return tailored

    def test_g2_approval_binds_plan_and_current_inputs(self):
        tailored = self.ready_plan()
        stored = self.service.save_tailoring_plan("BAT-RESEARCH", tailored)
        self.assertEqual(stored["status"], "READY_FOR_APPROVAL")
        with self.assertRaisesRegex(DomainError, "G2"):
            self.service.authorized_plan(tailored["plan_id"])
        with patch("career_v0.service.render_and_inspect") as render:
            with self.assertRaisesRegex(DomainError, "G2"):
                self.service.generate_resume_candidate(tailored["plan_id"])
            render.assert_not_called()
        path, export = self.service.export_plan_reviews("BAT-RESEARCH")
        wb = load_workbook(path)
        wb["方案总览"].cell(2, 1).value = "APPROVE"
        stream = io.BytesIO()
        wb.save(stream)
        preview = self.service.preview_plan_reviews(export["export_id"], stream.getvalue())
        with self.assertRaisesRegex(DomainError, "Owner action"):
            self.service.confirm_plan_reviews(preview["preview_id"], actor_channel="MODEL_OUTPUT")
        decision = self.service.confirm_plan_reviews(preview["preview_id"], actor_channel="SIMULATED_OWNER_TEST")
        self.assertEqual([row["decision"] for row in decision["applied"]], ["APPROVE"])
        self.assertEqual(self.service.authorized_plan(tailored["plan_id"])["plan"]["plan_id"], tailored["plan_id"])
        with patch.dict("os.environ", {"CAREER_NODE_BIN": shutil.which("node") or ""}):
            candidate = self.service.generate_resume_candidate(tailored["plan_id"])
        self.assertEqual(candidate["approval_decision_id"], decision["applied"][0]["decision_id"])
        self.assertEqual(len(candidate["content_project_ids"]), 7)
        self.assertEqual(candidate["status"], "NEEDS_PLAN_REVISION")
        self.assertTrue(Path(candidate["pdf_path"]).is_file())
        self.assertEqual(len(candidate["qa"]["page_image_paths"]), candidate["qa"]["page_count"])
        self.assertTrue(all(Path(image).is_file() for image in candidate["qa"]["page_image_paths"]))
        checks = {"both_pages_readable": True, "no_clipping_or_overlap": True,
                  "all_projects_visible": True}
        with self.assertRaisesRegex(DomainError, "Owner action"):
            self.service.review_resume_candidate(candidate["artifact_id"], decision="APPROVE",
                                                 reviewed_pdf_sha256=candidate["pdf_sha256"],
                                                 visual_checks=checks, owner_comment=None,
                                                 actor_channel="MODEL_OUTPUT")
        with self.assertRaisesRegex(DomainError, "failed two-page"):
            self.service.review_resume_candidate(candidate["artifact_id"], decision="APPROVE",
                                                 reviewed_pdf_sha256=candidate["pdf_sha256"],
                                                 visual_checks=checks, owner_comment=None,
                                                 actor_channel="SIMULATED_OWNER_TEST")
        with self.assertRaisesRegex(DomainError, "G3"):
            self.service.export_handoff(candidate["artifact_id"])
        selected = self.service.selected_jobs("BAT-RESEARCH")[0]
        self.service.save_research_dossier("BAT-RESEARCH", dossier(selected, full=True))
        with self.assertRaisesRegex(DomainError, "changed"):
            self.service.authorized_plan(tailored["plan_id"])

    def test_one_selected_job_has_one_formal_plan_id(self):
        tailored = self.ready_plan()
        self.service.save_tailoring_plan("BAT-RESEARCH", tailored)
        duplicate = copy.deepcopy(tailored)
        duplicate["plan_id"] = "PLAN-RESEARCH-OTHER"
        with self.assertRaisesRegex(DomainError, "one formal plan ID"):
            self.service.save_tailoring_plan("BAT-RESEARCH", duplicate)
        revised = copy.deepcopy(tailored)
        revised["plan_revision"] = 2
        revised["strategy"]["summary"] = "修订后的合成策略"
        self.assertEqual(self.service.save_tailoring_plan("BAT-RESEARCH", revised)["revision"], 2)

    def test_synthetic_g3_and_handoff_contract(self):
        """The renderer is mocked here; this proves G3/package state, not resume content QA."""
        tailored = self.ready_plan()
        self.service.save_tailoring_plan("BAT-RESEARCH", tailored)
        path, export = self.service.export_plan_reviews("BAT-RESEARCH")
        workbook = load_workbook(path)
        workbook["方案总览"].cell(2, 1).value = "APPROVE"
        preview = self.service.preview_plan_reviews(export["export_id"], save_workbook(workbook))
        self.service.confirm_plan_reviews(preview["preview_id"], actor_channel="SIMULATED_OWNER_TEST")

        def synthetic_pdf(_html_path, pdf_path, *, expected_project_ids, expected_text):
            pdf = FPDF(unit="pt", format="A4")
            for page_number in (1, 2):
                pdf.add_page()
                pdf.set_font("Helvetica", size=12)
                pdf.text(35, 50, f"Synthetic contract fixture page {page_number} with extractable resume text")
            pdf.output(str(pdf_path))
            images, hashes = [], []
            document = pdfium.PdfDocument(str(pdf_path))
            for index in range(2):
                image_path = pdf_path.with_name(f"{pdf_path.stem}-page-{index + 1}.png")
                page = document[index]
                page.render(scale=1).to_pil().save(image_path, format="PNG")
                page.close()
                images.append(str(image_path))
                hashes.append(hashlib.sha256(image_path.read_bytes()).hexdigest())
            document.close()
            return {"status": "WAITING_VISUAL_REVIEW", "problems": [], "page_count": 2,
                    "page_image_paths": images, "page_image_sha256": hashes,
                    "project_ids": expected_project_ids, "missing_text": [],
                    "hero_count": 4}

        with patch("career_v0.service.render_and_inspect", side_effect=synthetic_pdf):
            candidate = self.service.generate_resume_candidate(tailored["plan_id"])
        checks = {"both_pages_readable": True, "no_clipping_or_overlap": True,
                  "all_projects_visible": True}
        reviewed = self.service.review_resume_candidate(candidate["artifact_id"], decision="APPROVE",
                                                        reviewed_pdf_sha256=candidate["pdf_sha256"],
                                                        visual_checks=checks, owner_comment=None,
                                                        actor_channel="SIMULATED_OWNER_TEST")
        self.assertEqual(reviewed["status"], "OWNER_APPROVED")
        handoff = self.service.export_handoff(candidate["artifact_id"])
        self.assertEqual(handoff["status"], "HANDOFF_READY")
        self.assertEqual(handoff["application_status"], "NOT_FOUND")
        self.assertEqual(handoff["interview_research_status"], "NOT_STARTED")
        with zipfile.ZipFile(handoff["zip_path"]) as archive:
            self.assertEqual(set(archive.namelist()), {
                "manifest.json", "job_card.json", "application_card.json", "research.json",
                "full_jd.txt", "resume.pdf", "resume_readable.txt", "interview_research_prompt.txt",
            })
        self.assertEqual(self.service.export_handoff(candidate["artifact_id"])["handoff_id"],
                         handoff["handoff_id"])

    def test_planning_runner_uses_selected_full_jd_and_host_bound_refs(self):
        tailored = self.ready_plan()
        business = {
            "summary": tailored["strategy"]["summary"],
            "hero_project_ids": [row["project_id"] for row in tailored["project_plan"] if row["tier"] == "HERO"],
            "project_rationales": {row["project_id"]: row["rationale"] for row in tailored["project_plan"]},
            "project_emphasis": {row["project_id"]: row["emphasis"] for row in tailored["project_plan"]},
            "internship": {"action": "COMPRESS", "target_bullets": 2, "rationale": "保留实习事实并给项目留空间"},
            "risks": [],
        }
        trace = (json.dumps({"type": "thinking", "text": "private plan reasoning"}) + "\n" +
                 json.dumps({"type": "final", "text": json.dumps(business, ensure_ascii=False)},
                            ensure_ascii=False) + "\n").encode()
        dsh = Path(self.tmp.name) / "dsh.cmd"
        dsh.write_text("@echo off", encoding="utf-8")
        with patch("career_v0.planning_runner.subprocess.run",
                   return_value=CompletedProcess([], 0, stdout=trace, stderr=b"")) as invoke:
            result = prepare_tailoring_plan(self.service, "BAT-RESEARCH", tailored["job_ref"]["job_id"],
                                            dsh_bin=dsh, workspace=self.tmp.name)
        self.assertEqual(result["status"], "READY_FOR_APPROVAL")
        self.assertEqual(result["validation"]["project_count"], 7)
        self.assertIn(b"G1-selected job", invoke.call_args.kwargs["input"])
        self.assertNotIn(b"private plan reasoning", Path(result["planning_trace_path"]).read_bytes())
        with self.assertRaisesRegex(PlanningError, "REQUEST_CHANGES"):
            prepare_tailoring_plan(self.service, "BAT-RESEARCH", tailored["job_ref"]["job_id"],
                                   dsh_bin=dsh, workspace=self.tmp.name)

    def test_planning_runner_never_calls_model_before_g1(self):
        with self.assertRaisesRegex(DomainError, "G1"):
            prepare_tailoring_plan(self.service, "BAT-RESEARCH", self.job_ids[0],
                                   dsh_bin="missing")

    def test_selected_full_capture_bootstraps_honest_research(self):
        url = "https://example.org/job/1"
        discovery = self.service.ingest_discovery_trace("CMP-RESEARCH", event_trace())
        discovery = self.service.extend_discovery_trace(discovery["discovery_id"], full_fetch_trace(url))
        proposed = json.loads(json.loads(normalized_json().decode())["text"])
        proposed["source_url"] = url
        batch = self.service.create_batch("CMP-RESEARCH", [proposed], batch_id="BAT-CAPTURE",
                                          discovery_id=discovery["discovery_id"])
        self.assertEqual(batch["counts"]["qualified"], 1)
        path, export = self.service.export_candidates("BAT-CAPTURE")
        wb = load_workbook(path)
        wb["岗位候选"].cell(2, 1).value = "TRUE"
        preview = self.service.preview_selection(export["export_id"], save_workbook(wb))
        self.service.confirm_selection(preview["preview_id"], actor_channel="SIMULATED_OWNER_TEST",
                                       accept_shortfall=True)
        job_id = next(iter(batch["jobs"]))
        research = self.service.prepare_selected_source_research("BAT-CAPTURE", job_id)
        self.assertEqual(research["status"], "READY_FOR_PLAN")
        self.assertEqual(research["application_card"]["application_status"], "NEEDS_ORIGINAL_CHANNEL")
        self.assertEqual(research["business_scope_status"], "ROLE_JD_ONLY")
        self.assertIn("JD-R1", research["jd"]["requirement_quotes"])


def save_workbook(wb):
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


if __name__ == "__main__":
    unittest.main()
