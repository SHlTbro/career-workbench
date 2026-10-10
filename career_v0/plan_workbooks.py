"""Second V-0 XLSX round trip: per-plan review, with no approval on upload."""

from __future__ import annotations

import hashlib
import io
import uuid
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .domain import stable_hash, utc_now
from .plans import build_input_snapshot, input_snapshot_sha256, plan_content_sha256
from .workbooks import MAX_FILE_BYTES, WorkbookError, _readonly_key, _write_text_cell


SCHEMA = "plan-review-workbook@1"
EDITABLE = ("decision", "owner_comment", "owner_priority")
READ_ONLY = (
    "plan_id", "plan_revision", "job_id", "company", "title", "jd_url",
    "business_summary", "jd_focus", "tailoring_strategy", "hero_projects",
    "compact_projects", "protected_identity_plan", "highlights_edit",
    "skills_edit", "internship_edit", "order_change_summary", "template",
    "font_size", "page_target", "evidence_summary", "risk_summary",
    "plan_detail_ref", "input_snapshot_sha256", "plan_content_sha256",
)
COLUMNS = EDITABLE + READ_ONLY
PROJECT_COLUMNS = ("plan_id", "project_id", "source_order", "tier", "target_bullets", "target_lines", "purpose", "core_stack", "rationale")
SOURCE_COLUMNS = ("plan_id", "business_evidence_refs", "jd_requirement_refs", "fde_map_refs", "risk_summary")
DECISIONS = {"PENDING", "APPROVE", "REQUEST_CHANGES", "SKIP"}


def _text(values: Any) -> str:
    if isinstance(values, list):
        return ", ".join(str(value) for value in values)
    if values is None:
        return ""
    return str(values)


def _projection(plan: dict[str, Any], job: dict[str, Any], research: dict[str, Any]) -> dict[str, Any]:
    strategy = plan["strategy"]
    projects = plan["project_plan"]
    signals = plan["protected_signal_plan"]
    layout = plan["layout_plan"]
    internship = next((row for row in plan["section_plan"] if row["section_id"] == "internship"), {})
    return {
        "plan_id": plan["plan_id"], "plan_revision": plan["plan_revision"],
        "job_id": plan["job_ref"]["job_id"], "company": job.get("company"),
        "title": job.get("title"), "jd_url": research.get("jd", {}).get("source_url"),
        "business_summary": research.get("business_summary"),
        "jd_focus": _text(strategy.get("jd_requirement_refs")),
        "tailoring_strategy": strategy.get("summary"),
        "hero_projects": _text([row["project_id"] for row in projects if row["tier"] == "HERO"]),
        "compact_projects": _text([row["project_id"] for row in projects if row["tier"] == "COMPACT"]),
        "protected_identity_plan": _text([f"{key}: {value.get('handling', '')}" for key, value in sorted(signals.items())]),
        "highlights_edit": _text(strategy.get("highlights_edit")),
        "skills_edit": _text(strategy.get("skills_edit")),
        "internship_edit": f"{internship.get('action', '')}: {internship.get('rationale', '')}",
        "order_change_summary": _text([row.get("rationale") for row in plan.get("order_change_proposal", [])]) or "保持原顺序",
        "template": layout.get("template_id"), "font_size": layout.get("font_size_pt"),
        "page_target": layout.get("target_pages"),
        "evidence_summary": _text(strategy.get("business_evidence_refs")) + " | " + _text(strategy.get("fde_map_refs")),
        "risk_summary": _text(plan.get("risks")),
        "plan_detail_ref": f"{plan['plan_id']}@{plan['plan_revision']}",
        "input_snapshot_sha256": input_snapshot_sha256(build_input_snapshot(plan)),
        "plan_content_sha256": plan_content_sha256(plan),
    }


def _sheet_rows(rows: list[dict[str, Any]], columns: tuple[str, ...]) -> dict[str, dict[str, Any]]:
    return {row["plan_id"] + (":" + row["project_id"] if "project_id" in row else ""): {key: row.get(key) for key in columns} for row in rows}


def export_plan_workbook(campaign_id: str, batch_id: str, entries: list[dict[str, Any]],
                         inventory: list[dict[str, Any]]) -> tuple[bytes, dict[str, Any]]:
    if not entries:
        raise WorkbookError("no plans to review")
    export_id = "PEXP-" + uuid.uuid4().hex[:20].upper()
    rows: dict[str, dict[str, Any]] = {}
    project_rows: list[dict[str, Any]] = []
    source_rows: list[dict[str, Any]] = []
    inventory_by_id = {item["project_id"]: item for item in inventory}
    for entry in entries:
        plan, job, research = entry["plan"], entry["job"], entry["research"]
        if plan["campaign_id"] != campaign_id or job["job_id"] != plan["job_ref"]["job_id"]:
            raise WorkbookError("plan/campaign/job identity mismatch")
        plan_id = plan["plan_id"]
        if plan_id in rows:
            raise WorkbookError("duplicate plan_id")
        rows[plan_id] = _projection(plan, job, research)
        for row in plan["project_plan"]:
            project = inventory_by_id[row["project_id"]]
            project_rows.append({"plan_id": plan_id, "project_id": row["project_id"],
                                 "source_order": row["source_order"], "tier": row["tier"],
                                 "target_bullets": row.get("target_bullets"), "target_lines": row.get("target_lines"),
                                 "purpose": project.get("purpose"), "core_stack": _text(project.get("core_stack")),
                                 "rationale": row.get("rationale")})
        source_rows.append({"plan_id": plan_id,
                            "business_evidence_refs": _text(plan["strategy"].get("business_evidence_refs")),
                            "jd_requirement_refs": _text(plan["strategy"].get("jd_requirement_refs")),
                            "fde_map_refs": _text(plan["strategy"].get("fde_map_refs")),
                            "risk_summary": _text(plan.get("risks"))})
    snapshot = {"schema": SCHEMA, "export_id": export_id, "campaign_id": campaign_id,
                "batch_id": batch_id, "created_at": utc_now(), "rows": rows,
                "project_rows": _sheet_rows(project_rows, PROJECT_COLUMNS),
                "source_rows": _sheet_rows(source_rows, SOURCE_COLUMNS)}
    snapshot["snapshot_sha256"] = stable_hash({key: snapshot[key] for key in ("campaign_id", "batch_id", "rows", "project_rows", "source_rows")})
    wb = Workbook()
    instructions = wb.active
    instructions.title = "使用说明"
    for row in (("Career Agent V-0 逐岗位方案预审",),
                ("可编辑", "只编辑方案总览的 decision、owner_comment、owner_priority。"),
                ("决策", "PENDING / APPROVE / REQUEST_CHANGES / SKIP；默认 PENDING。"),
                ("批准", "上传后仍须在 Career 工作台逐行预览并确认；APPROVE 与实质修改意见同时出现时不得直接批准。"),
                ("版本", "每行绑定方案和输入版本；过期行单独报冲突，不影响其他有效行。")):
        instructions.append(row)
    instructions.column_dimensions["A"].width = 24
    instructions.column_dimensions["B"].width = 110
    main = wb.create_sheet("方案总览")
    main.append(COLUMNS)
    for plan_id in sorted(rows):
        row = ["PENDING", None, None] + [rows[plan_id][key] for key in READ_ONLY]
        main.append(row)
        for col, value in enumerate(row, 1):
            _write_text_cell(main.cell(main.max_row, col), value)
    main.freeze_panes = "G2"
    main.auto_filter.ref = f"A1:{main.cell(1, len(COLUMNS)).column_letter}{max(main.max_row, 2)}"
    validation = DataValidation(type="list", formula1='"PENDING,APPROVE,REQUEST_CHANGES,SKIP"', allow_blank=False)
    validation.showErrorMessage = True
    main.add_data_validation(validation)
    validation.add(f"A2:A{max(main.max_row, 2)}")
    for index, field in enumerate(COLUMNS, 1):
        cell = main.cell(1, index)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2763A6" if field in EDITABLE else "4C566A")
        main.column_dimensions[cell.column_letter].width = 25 if field in EDITABLE else 29
    for name, columns, values in (("逐项目取舍", PROJECT_COLUMNS, project_rows), ("来源与差距", SOURCE_COLUMNS, source_rows)):
        sheet = wb.create_sheet(name)
        sheet.append(columns)
        for record in values:
            sheet.append([record.get(key) for key in columns])
            for cell in sheet[sheet.max_row]:
                _write_text_cell(cell, cell.value)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{sheet.cell(1, len(columns)).column_letter}{max(sheet.max_row, 2)}"
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 28
    meta = wb.create_sheet("_meta")
    for key in ("schema", "export_id", "campaign_id", "batch_id", "snapshot_sha256"):
        meta.append((key, snapshot[key]))
    meta.sheet_state = "hidden"
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue(), snapshot


def _read_table(wb: Any, name: str, columns: tuple[str, ...], key_fields: tuple[str, ...]) -> dict[str, dict[str, Any]]:
    if name not in wb.sheetnames:
        raise WorkbookError(f"{name} sheet missing")
    sheet = wb[name]
    if [cell.value for cell in sheet[1]] != list(columns):
        raise WorkbookError(f"{name} headers changed")
    rows: dict[str, dict[str, Any]] = {}
    for row_number, cells in enumerate(sheet.iter_rows(min_row=2, max_col=len(columns)), 2):
        if all(cell.value is None for cell in cells):
            continue
        record = {field: cells[index].value for index, field in enumerate(columns)}
        key = ":".join(str(record[field]) for field in key_fields)
        if key in rows:
            raise WorkbookError(f"{name} duplicate identity at row {row_number}")
        rows[key] = record
    return rows


def _compare_readonly(actual: dict[str, dict[str, Any]], expected: dict[str, dict[str, Any]],
                      columns: tuple[str, ...], name: str) -> None:
    if set(actual) != set(expected):
        raise WorkbookError(f"{name} plan/project identity set changed")
    for key, row in actual.items():
        for field in columns:
            if _readonly_key(row[field]) != _readonly_key(expected[key][field]):
                raise WorkbookError(f"{name} {key}: read-only conflict in {field}")


def import_plan_workbook(data: bytes, snapshot: dict[str, Any],
                         current_bindings: dict[str, dict[str, Any]]) -> dict[str, Any]:
    if not isinstance(data, bytes) or not data or len(data) > MAX_FILE_BYTES:
        raise WorkbookError("plan workbook missing or larger than 10 MB")
    try:
        wb = load_workbook(io.BytesIO(data), data_only=False, keep_links=False)
    except Exception as exc:
        raise WorkbookError("plan workbook is not readable XLSX") from exc
    for tab in wb:
        if any(cell.data_type == "f" for row in tab for cell in row):
            raise WorkbookError(f"formula not allowed in {tab.title}")
    meta_rows = list(wb["_meta"].values) if "_meta" in wb.sheetnames else []
    if len(meta_rows) != len({row[0] for row in meta_rows if row and row[0] is not None}):
        raise WorkbookError("duplicate _meta key")
    meta = {row[0]: row[1] for row in meta_rows if len(row) >= 2}
    for key in ("schema", "export_id", "campaign_id", "batch_id", "snapshot_sha256"):
        if meta.get(key) != snapshot[key]:
            raise WorkbookError(f"_meta conflict: {key}")
    main = _read_table(wb, "方案总览", COLUMNS, ("plan_id",))
    _compare_readonly(main, snapshot["rows"], READ_ONLY, "方案总览")
    projects = _read_table(wb, "逐项目取舍", PROJECT_COLUMNS, ("plan_id", "project_id"))
    _compare_readonly(projects, snapshot["project_rows"], PROJECT_COLUMNS, "逐项目取舍")
    sources = _read_table(wb, "来源与差距", SOURCE_COLUMNS, ("plan_id",))
    _compare_readonly(sources, snapshot["source_rows"], SOURCE_COLUMNS, "来源与差距")
    valid: list[dict[str, Any]] = []
    conflicts: list[dict[str, Any]] = []
    for plan_id, row in sorted(main.items()):
        decision = row["decision"] or "PENDING"
        if decision not in DECISIONS:
            raise WorkbookError(f"{plan_id}: invalid decision")
        comment = row["owner_comment"]
        if comment is not None and (not isinstance(comment, str) or len(comment) > 2000):
            raise WorkbookError(f"{plan_id}: invalid owner_comment")
        priority = row["owner_priority"]
        if priority not in (None, "", "P0", "P1", "P2"):
            raise WorkbookError(f"{plan_id}: invalid owner_priority")
        expected = snapshot["rows"][plan_id]
        current = current_bindings.get(plan_id)
        if current is None or any(current.get(key) != expected[key] for key in ("plan_revision", "input_snapshot_sha256", "plan_content_sha256")):
            conflicts.append({"plan_id": plan_id, "reason": "STALE_PLAN_OR_INPUT"})
            continue
        if decision == "APPROVE" and comment and comment.strip():
            conflicts.append({"plan_id": plan_id, "reason": "APPROVAL_WITH_CHANGE_REQUEST"})
            continue
        if decision == "REQUEST_CHANGES" and not (comment and comment.strip()):
            conflicts.append({"plan_id": plan_id, "reason": "CHANGE_REQUEST_NEEDS_COMMENT"})
            continue
        valid.append({"plan_id": plan_id, "plan_revision": expected["plan_revision"],
                      "input_snapshot_sha256": expected["input_snapshot_sha256"],
                      "plan_content_sha256": expected["plan_content_sha256"],
                      "decision": decision, "owner_comment": comment or None,
                      "owner_priority": priority or None})
    return {"schema": "plan-import-preview@1", "preview_id": "PIMP-" + uuid.uuid4().hex[:20].upper(),
            "export_id": snapshot["export_id"], "campaign_id": snapshot["campaign_id"],
            "batch_id": snapshot["batch_id"], "file_sha256": hashlib.sha256(data).hexdigest(),
            "valid_rows": valid, "conflicts": conflicts,
            "status": "PENDING_CONFIRMATION", "created_at": utc_now()}
