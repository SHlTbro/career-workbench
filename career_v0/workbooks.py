"""Candidate XLSX exchange contract. A workbook never authorizes research."""

from __future__ import annotations

import hashlib
import io
import uuid
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .domain import stable_hash, utc_now


class WorkbookError(ValueError):
    pass


SCHEMA = "candidate-workbook@1"
EDITABLE = ("selected", "owner_priority", "owner_note")
READ_ONLY = (
    "job_id", "company", "title", "recruitment_type", "graduation_eligibility",
    "city", "salary_raw", "salary_min_monthly", "salary_basis", "salary_check",
    "eligibility_check", "scope_check", "opening_status", "jd_url", "apply_url", "application_status",
    "jd_completeness", "source", "published_at", "captured_at",
    "match_summary", "key_tags", "risk_flags",
)
COLUMNS = EDITABLE + READ_ONLY
VIEW_COLUMNS = ("job_id", "company", "title", "city", "salary_raw", "source_url", "risk_flags")
MAX_FILE_BYTES = 10 * 1024 * 1024


def _cell_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (list, dict)):
        return ", ".join(map(str, value)) if isinstance(value, list) else stable_hash(value)
    return str(value)


def _write_text_cell(cell: Any, value: Any) -> None:
    cell.value = _cell_value(value)
    if isinstance(cell.value, str):
        # XLSX cell type s keeps source strings beginning = + - @ as data.
        cell.data_type = "s"
        cell.number_format = "@"


def _job_projection(job: dict[str, Any]) -> dict[str, Any]:
    result = {field: _cell_value(job.get(field)) for field in READ_ONLY}
    result["graduation_eligibility"] = result["graduation_eligibility"] or job.get("graduation_years")
    result["source"] = result["source"] or job.get("source_platform")
    result["jd_url"] = result["jd_url"] or job.get("source_url")
    return {key: _cell_value(value) for key, value in result.items()}


def _readonly_key(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return str(value).removesuffix(".0")
    return str(value).strip()


def export_candidate_workbook(campaign: dict[str, Any], batch: dict[str, Any]) -> tuple[bytes, dict[str, Any]]:
    if campaign.get("status") != "PREFERENCE_CONFIRMED":
        raise WorkbookError("G0 confirmation required")
    if batch.get("campaign_id") != campaign.get("campaign_id") or batch.get("campaign_revision") != campaign.get("revision"):
        raise WorkbookError("batch/campaign version mismatch")
    export_id = "EXP-" + uuid.uuid4().hex[:20].upper()
    qualified = sorted((j for j in batch["jobs"].values() if j["pool"] == "qualified"), key=lambda j: j["job_id"])
    unknown = sorted((j for j in batch["jobs"].values() if j["pool"] == "unknown"), key=lambda j: j["job_id"])
    excluded = sorted((j for j in batch["jobs"].values() if j["pool"] == "excluded"), key=lambda j: j["job_id"])
    snapshot_rows = {job["job_id"]: _job_projection(job) for job in qualified}
    snapshot_views = {
        name: {job["job_id"]: {field: _cell_value(job.get(field)) for field in VIEW_COLUMNS} for job in records}
        for name, records in (("待核实", unknown), ("排除摘要", excluded))
    }
    snapshot = {
        "schema": SCHEMA, "export_id": export_id,
        "campaign_id": campaign["campaign_id"], "campaign_revision": campaign["revision"],
        "batch_id": batch["batch_id"], "batch_revision": batch["revision"],
        "created_at": utc_now(), "rows": snapshot_rows, "views": snapshot_views,
        "selection_mode": "TRUE_FALSE_DROPDOWN",
    }
    snapshot["snapshot_sha256"] = stable_hash({k: snapshot[k] for k in ("campaign_id", "campaign_revision", "batch_id", "batch_revision", "rows", "views")})

    wb = Workbook()
    instructions = wb.active
    instructions.title = "使用说明"
    for row in (
        ("Career Agent V-0 岗位候选表",),
        ("招聘方向", "2027 届秋招"),
        ("Campaign", campaign["campaign_id"]),
        ("导出时间 UTC", snapshot["created_at"]),
        ("操作", "只编辑 岗位候选 Sheet 的 selected、owner_priority、owner_note。"),
        ("选择", "selected 默认为 FALSE；使用 TRUE/FALSE 下拉。若客户端支持原生单元格复选框，可将其绑定到布尔值。"),
        ("确认", "上传文件后仍需在 Career 工作台检查预览并点击“确认导入并研究”。"),
        ("数量", f"合格 {len(qualified)}；待核实 {len(unknown)}；排除 {len(excluded)}；目标 {batch['target_count']}。"),
        ("未知", "待核实岗位不算已满足薪资/届别条件；没有 50 条时不补假岗位。"),
    ):
        instructions.append(row)
    instructions.column_dimensions["A"].width = 24
    instructions.column_dimensions["B"].width = 105

    sheet = wb.create_sheet("岗位候选")
    sheet.append(COLUMNS)
    for job in qualified:
        projection = snapshot_rows[job["job_id"]]
        row = [False, None, None] + [projection[field] for field in READ_ONLY]
        sheet.append(row)
        row_number = sheet.max_row
        for col, value in enumerate(row, 1):
            _write_text_cell(sheet.cell(row_number, col), value)
    sheet.freeze_panes = "E2"
    sheet.auto_filter.ref = f"A1:{sheet.cell(1, len(COLUMNS)).column_letter}{max(sheet.max_row, 2)}"
    validation = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=True)
    validation.error = "请选择 TRUE 或 FALSE"
    validation.errorTitle = "选择值无效"
    validation.showErrorMessage = True
    sheet.add_data_validation(validation)
    validation.add(f"A2:A{max(sheet.max_row, 2)}")
    for col, field in enumerate(COLUMNS, 1):
        cell = sheet.cell(1, col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2763A6" if field in EDITABLE else "4C566A")
        cell.alignment = Alignment(wrap_text=True)
        sheet.column_dimensions[cell.column_letter].width = 18 if field in EDITABLE else (34 if field in {"company", "title", "jd_url", "apply_url", "match_summary"} else 22)
    for row in sheet.iter_rows(min_row=2, max_col=3):
        for cell in row:
            cell.fill = PatternFill("solid", fgColor="E7F3FF")

    for name, records in (("待核实", unknown), ("排除摘要", excluded)):
        view = wb.create_sheet(name)
        view.append(VIEW_COLUMNS)
        for job in records:
            values = [job.get(key) for key in VIEW_COLUMNS]
            view.append([_cell_value(value) for value in values])
            for cell in view[view.max_row]:
                _write_text_cell(cell, cell.value)
        view.freeze_panes = "A2"
        view.auto_filter.ref = f"A1:G{max(view.max_row, 2)}"
        for letter in "ABCDEFG":
            view.column_dimensions[letter].width = 26

    meta = wb.create_sheet("_meta")
    for key in ("schema", "export_id", "campaign_id", "campaign_revision", "batch_id", "batch_revision", "snapshot_sha256", "selection_mode"):
        meta.append((key, snapshot[key]))
    meta.sheet_state = "hidden"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue(), snapshot


def _read_meta(wb: Any) -> dict[str, Any]:
    if "_meta" not in wb.sheetnames:
        raise WorkbookError("_meta sheet missing")
    rows = list(wb["_meta"].values)
    keys = [str(row[0]) for row in rows if row and row[0] is not None]
    if len(keys) != len(set(keys)):
        raise WorkbookError("duplicate _meta key")
    meta = {str(k): v for k, v, *_ in rows if k is not None}
    return meta


def _check_view(wb: Any, name: str, expected: dict[str, dict[str, Any]]) -> None:
    if name not in wb.sheetnames:
        raise WorkbookError(f"{name} sheet missing")
    sheet = wb[name]
    headers = [cell.value for cell in sheet[1]]
    if headers != list(VIEW_COLUMNS):
        raise WorkbookError(f"{name} headers changed")
    seen: set[str] = set()
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, max_col=len(VIEW_COLUMNS)), 2):
        if all(cell.value is None for cell in row):
            continue
        job_id = row[0].value
        if job_id not in expected or job_id in seen:
            raise WorkbookError(f"{name} row {row_number}: unknown or duplicate job_id")
        seen.add(job_id)
        for index, field in enumerate(VIEW_COLUMNS):
            if _readonly_key(row[index].value) != _readonly_key(expected[job_id][field]):
                raise WorkbookError(f"{name} row {row_number}: read-only conflict in {field}")
    if seen != set(expected):
        raise WorkbookError(f"{name} missing job_id rows")


def _selected(value: Any, row_number: int) -> bool:
    if value is None or value == "":
        return False
    if value is True or value == "TRUE":
        return True
    if value is False or value == "FALSE":
        return False
    raise WorkbookError(f"岗位候选 row {row_number}: selected must be boolean TRUE/FALSE")


def import_candidate_workbook(data: bytes, snapshot: dict[str, Any], *, previous_selected: list[str] | None = None) -> dict[str, Any]:
    if not isinstance(data, bytes) or len(data) > MAX_FILE_BYTES:
        raise WorkbookError("candidate workbook missing or larger than 10 MB")
    try:
        wb = load_workbook(io.BytesIO(data), data_only=False, read_only=False, keep_links=False)
    except Exception as exc:
        raise WorkbookError("candidate workbook is not a readable XLSX") from exc
    for tab in wb:
        if any(cell.data_type == "f" for row in tab for cell in row):
            raise WorkbookError(f"formula not allowed in {tab.title}")
    meta = _read_meta(wb)
    for key in ("schema", "export_id", "campaign_id", "campaign_revision", "batch_id", "batch_revision", "snapshot_sha256"):
        if meta.get(key) != snapshot.get(key):
            raise WorkbookError(f"_meta conflict: {key}")
    if "岗位候选" not in wb.sheetnames:
        raise WorkbookError("岗位候选 sheet missing")
    for name, expected in snapshot["views"].items():
        _check_view(wb, name, expected)
    sheet = wb["岗位候选"]
    headers = [cell.value for cell in sheet[1]]
    if len(headers) != len(set(headers)) or any(key not in headers for key in COLUMNS):
        raise WorkbookError("岗位候选 headers missing or duplicated")
    positions = {key: headers.index(key) for key in COLUMNS}
    selected: list[str] = []
    priorities: dict[str, str | None] = {}
    notes: dict[str, str | None] = {}
    seen: set[str] = set()
    for row_number, row in enumerate(sheet.iter_rows(min_row=2), 2):
        if all(cell.value is None for cell in row):
            continue
        if any(cell.data_type == "f" for cell in row):
            raise WorkbookError(f"岗位候选 row {row_number}: formula not allowed")
        job_id = row[positions["job_id"]].value
        if job_id not in snapshot["rows"]:
            raise WorkbookError(f"岗位候选 row {row_number}: unknown job_id {job_id}")
        if job_id in seen:
            raise WorkbookError(f"岗位候选 row {row_number}: duplicate job_id {job_id}")
        seen.add(job_id)
        expected = snapshot["rows"][job_id]
        for field in READ_ONLY:
            actual = row[positions[field]].value
            if _readonly_key(actual) != _readonly_key(expected.get(field)):
                raise WorkbookError(f"岗位候选 row {row_number} job {job_id}: read-only conflict in {field}")
        if _selected(row[positions["selected"]].value, row_number):
            selected.append(job_id)
        priority = row[positions["owner_priority"]].value
        if priority not in (None, "", "P0", "P1", "P2"):
            raise WorkbookError(f"岗位候选 row {row_number}: invalid owner_priority")
        priorities[job_id] = priority or None
        note = row[positions["owner_note"]].value
        if note is not None and (not isinstance(note, str) or len(note) > 2000):
            raise WorkbookError(f"岗位候选 row {row_number}: invalid owner_note")
        notes[job_id] = note or None
    missing = set(snapshot["rows"]) - seen
    if missing:
        raise WorkbookError(f"岗位候选 missing job_id rows: {', '.join(sorted(missing)[:5])}")
    before = set(previous_selected or [])
    after = set(selected)
    return {
        "schema": "candidate-import-preview@1", "preview_id": "IMP-" + uuid.uuid4().hex[:20].upper(),
        "export_id": snapshot["export_id"], "campaign_id": snapshot["campaign_id"],
        "campaign_revision": snapshot["campaign_revision"], "batch_id": snapshot["batch_id"],
        "batch_revision": snapshot["batch_revision"], "file_sha256": hashlib.sha256(data).hexdigest(),
        "selected_job_ids": sorted(after), "added_job_ids": sorted(after - before),
        "removed_job_ids": sorted(before - after), "owner_priority": priorities,
        "owner_note": notes, "created_at": utc_now(), "status": "PENDING_CONFIRMATION",
    }
